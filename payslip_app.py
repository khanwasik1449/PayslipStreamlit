import streamlit as st
from docx import Document
from docx.shared import Pt
from docx2pdf import convert
import os
from datetime import datetime

# Folder to save generated payslips
SAVE_FOLDER = "Payslips"
os.makedirs(SAVE_FOLDER, exist_ok=True)

st.title("💼 Payslip Generator (Fixed Template)")

# Employee Input
employeeName = st.text_input("Employee Name")
pin = st.text_input("PIN")
designation = st.text_input("Designation")
joiningDate = st.date_input("Joining Date")
salary = st.number_input("Basic Salary", value=0)
allowance = st.number_input("Allowance", value=0)
transport = st.number_input("Transport Deduction", value=0)
tax = st.number_input("Tax Deduction", value=0)
otherDeductions = st.number_input("Other Deductions", value=0)
payslipMonth = st.text_input("Payslip Month (e.g., August 2025)")

if st.button("Generate Payslip"):
    if not all([employeeName, pin, designation, payslipMonth]):
        st.error("Please fill all required fields")
    else:
        # Salary calculations
        basic_comp = round(0.5 * salary)
        house_rent = round(0.3 * salary)
        medical = round(0.1 * salary)
        conveyance = round(0.1 * salary)
        totalDeductions = transport + tax + otherDeductions
        netSalary = salary + allowance - totalDeductions

        # Load fixed template
        TEMPLATE_PATH = "Payslip_Template.docx"  # put your template in the same folder
        doc = Document(TEMPLATE_PATH)

        # Mapping placeholders to actual values
        placeholders = {
            "{{employeeName}}": employeeName,
            "{{pin}}": pin,
            "{{designation}}": designation,
            "{{joiningDate}}": str(joiningDate),
            "{{salary}}": f"BDT {salary:,.2f}",
            "{{allowance}}": f"BDT {allowance:,.2f}",
            "{{transport}}": f"BDT {transport:,.2f}",
            "{{tax}}": f"BDT {tax:,.2f}",
            "{{otherDeductions}}": f"BDT {otherDeductions:,.2f}",
            "{{totalDeductions}}": f"BDT {totalDeductions:,.2f}",
            "{{netSalary}}": f"BDT {netSalary:,.2f}",
            "{{basic}}": f"BDT {basic_comp:,.2f}",
            "{{houseRent}}": f"BDT {house_rent:,.2f}",
            "{{medical}}": f"BDT {medical:,.2f}",
            "{{conveyance}}": f"BDT {conveyance:,.2f}",
            "{{payslipMonth}}": payslipMonth
        }

        # Replace placeholders
        for p in doc.paragraphs:
            for key, val in placeholders.items():
                if key in p.text:
                    p.text = p.text.replace(key, val)
                    for run in p.runs:
                        run.font.size = Pt(12)

        # Save DOCX
        filename_docx = f"Payslip_{employeeName}_{payslipMonth}.docx"
        filepath_docx = os.path.join(SAVE_FOLDER, filename_docx)
        doc.save(filepath_docx)

        # Convert to PDF
        filename_pdf = filepath_docx.replace(".docx", ".pdf")
        convert(filepath_docx, filename_pdf)

        st.success("Payslip generated successfully! 🎉")
        with open(filename_pdf, "rb") as f:
            st.download_button("💾 Download Payslip PDF", f, file_name=os.path.basename(filename_pdf))
